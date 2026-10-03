"""PDF -> Excel pipeline.

    PDF -> has text?  YES -> Camelot / pdfplumber
                      NO  -> OCR (Tesseract + img2table)
        -> detect tables -> clean data -> pandas -> Excel (.xlsx)
"""
from __future__ import annotations

import logging
import os
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import pdfplumber

log = logging.getLogger(__name__)

MIN_TEXT_CHARS = 25
OCR_LANG = os.environ.get("OCR_LANG", "eng")


def _ensure_tesseract_on_path() -> bool:
    if shutil.which("tesseract"):
        return True
    candidates = [
        os.environ.get("TESSERACT_DIR"),
        r"C:\Program Files\Tesseract-OCR",
        r"C:\Program Files (x86)\Tesseract-OCR",
        os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR"),
    ]
    for folder in candidates:
        if folder and Path(folder, "tesseract.exe").exists():
            os.environ["PATH"] = folder + os.pathsep + os.environ.get("PATH", "")
            return True
    return False


TESSERACT_AVAILABLE = _ensure_tesseract_on_path()


SHEET_NAME = "Tables"


@dataclass
class TableResult:
    page: int
    method: str
    df: pd.DataFrame


@dataclass
class TableBlock:
    """One table on the output sheet; a table continued over pages is a single block."""
    pages: list[int]
    methods: list[str]
    df: pd.DataFrame
    start_row: int = 0


@dataclass
class PageReport:
    page: int
    has_text: bool
    method: str
    tables: int


@dataclass
class ConversionResult:
    tables: list[TableResult] = field(default_factory=list)
    blocks: list[TableBlock] = field(default_factory=list)
    pages: list[PageReport] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# 1. Does the page contain text?
# --------------------------------------------------------------------------- #
def classify_pages(pdf_path: Path) -> dict[int, bool]:
    with pdfplumber.open(pdf_path) as pdf:
        return {
            number: len((page.extract_text() or "").strip()) >= MIN_TEXT_CHARS
            for number, page in enumerate(pdf.pages, start=1)
        }


# --------------------------------------------------------------------------- #
# 2a. Text pages: ruled rows -> Camelot (lattice) -> pdfplumber -> Camelot (stream)
# --------------------------------------------------------------------------- #
def _camelot_tables(pdf_path: Path, pages: list[int], flavor: str) -> dict[int, list[pd.DataFrame]]:
    import camelot

    found: dict[int, list[pd.DataFrame]] = {}
    try:
        tables = camelot.read_pdf(
            pdf_path.read_bytes(),  # a path would stay open (lattice), blocking job cleanup on Windows
            pages=",".join(map(str, pages)),
            flavor=flavor,
            suppress_stdout=True,
        )
    except Exception as exc:  # camelot raises on some malformed pages
        log.warning("Camelot %s failed: %s", flavor, exc)
        return found
    for table in tables:
        found.setdefault(int(table.page), []).append(table.df)
    return found


def _pdfplumber_tables(pdf_path: Path, pages: list[int]) -> dict[int, list[pd.DataFrame]]:
    found: dict[int, list[pd.DataFrame]] = {}
    with pdfplumber.open(pdf_path) as pdf:
        for number in pages:
            for raw in pdf.pages[number - 1].extract_tables():
                if raw:
                    found.setdefault(number, []).append(pd.DataFrame(raw))
    return found


RULE_MAX_THICKNESS = 2.0
SAME_Y = 1.5
SAME_X = 3.0
MAX_ROW_HEIGHT = 150.0


def _rule_rows(page) -> list[tuple[float, list[tuple[float, float]]]]:
    """Horizontal rules (thin rects / lines) grouped by their y position."""
    segments = [
        (r["x0"], r["x1"], r["top"])
        for r in page.rects
        if r["height"] <= RULE_MAX_THICKNESS and r["width"] > SAME_X
    ] + [
        (l["x0"], l["x1"], l["top"])
        for l in page.lines
        if abs(l["bottom"] - l["top"]) <= RULE_MAX_THICKNESS and l["x1"] - l["x0"] > SAME_X
    ]
    rows: list[tuple[float, list[tuple[float, float]]]] = []
    for x0, x1, y in sorted(segments, key=lambda s: s[2]):
        if rows and y - rows[-1][0] <= SAME_Y:
            rows[-1][1].append((x0, x1))
        else:
            rows.append((y, [(x0, x1)]))
    return rows


def _joints(segments: list[tuple[float, float]]) -> list[float]:
    """x positions where one rule segment ends and the next begins (column borders)."""
    joints, end = [], None
    for x0, x1 in sorted(segments):
        if end is not None and x0 >= end - SAME_Y:
            joints.append((end + x0) / 2)
        end = x1 if end is None else max(end, x1)
    return joints


def _find_ruled_grids(page) -> list[tuple[list[float], list[float]]]:
    """Tables whose rows are separated by horizontal rules drawn as one segment
    per column. Rows come from the rules, columns from the segment joints, so
    text outside the ruled area (letterhead, totals, footers) is never included."""
    groups: list[list[tuple[float, float, float, list[float]]]] = []
    for y, segments in _rule_rows(page):
        rule = (y, min(s[0] for s in segments), max(s[1] for s in segments), _joints(segments))
        prev = groups[-1][-1] if groups else None
        if (
            prev
            and abs(rule[1] - prev[1]) <= SAME_X
            and abs(rule[2] - prev[2]) <= SAME_X
            and y - prev[0] <= MAX_ROW_HEIGHT
        ):
            groups[-1].append(rule)
        else:
            groups.append([rule])

    grids = []
    for group in groups:
        # Rules without column joints at the edges are underlines / separators, not the table.
        while group and not group[0][3]:
            group.pop(0)
        while group and not group[-1][3]:
            group.pop()
        if len(group) < 2:
            continue
        clusters: list[list[float]] = []
        for x in sorted(x for rule in group for x in rule[3]):
            if clusters and x - clusters[-1][-1] <= SAME_X:
                clusters[-1].append(x)
            else:
                clusters.append([x])
        borders = [sum(c) / len(c) for c in clusters if len(c) >= len(group) / 2]
        if not borders:
            continue
        left = min(rule[1] for rule in group)
        right = max(rule[2] for rule in group)
        xs = [left, *borders, right]
        ys = _add_unruled_edge_rows(page, xs, [rule[0] for rule in group])
        if len(ys) >= 3:  # at least two rows
            grids.append((xs, ys))
    return grids


def _add_unruled_edge_rows(page, xs: list[float], ys: list[float]) -> list[float]:
    """A table continued from the previous page often has no rule above its
    first row (or below its last). Add that row when its words sit inside
    the columns like a table row, so text crossing columns is never added."""
    heights = sorted(b - a for a, b in zip(ys, ys[1:]))
    reach = 1.25 * heights[len(heights) // 2]
    words = [w for w in page.extract_words() if w["x1"] > xs[0] and w["x0"] < xs[-1]]

    def looks_like_row(band: list[dict]) -> bool:
        columns = set()
        for w in band:
            column = next(
                (i for i in range(len(xs) - 1) if xs[i] - 1 <= w["x0"] and w["x1"] <= xs[i + 1] + 1), None
            )
            if column is None:
                return False
            columns.add(column)
        return len(columns) >= max(2, (len(xs) - 1) / 2)

    def center(w: dict) -> float:
        return (w["top"] + w["bottom"]) / 2

    above = [w for w in words if ys[0] - reach <= center(w) < ys[0]]
    if above and looks_like_row(above):
        ys = [min(w["top"] for w in above) - 1, *ys]
    below = [w for w in words if ys[-1] < center(w) <= ys[-1] + reach]
    if below and looks_like_row(below):
        ys = [*ys, max(w["bottom"] for w in below) + 1]
    return ys


def _ruled_tables(pdf_path: Path, pages: list[int]) -> dict[int, list[pd.DataFrame]]:
    found: dict[int, list[pd.DataFrame]] = {}
    with pdfplumber.open(pdf_path) as pdf:
        for number in pages:
            page = pdf.pages[number - 1]
            for xs, ys in _find_ruled_grids(page):
                rows = page.extract_table({
                    "vertical_strategy": "explicit",
                    "horizontal_strategy": "explicit",
                    "explicit_vertical_lines": xs,
                    "explicit_horizontal_lines": ys,
                    "text_x_tolerance": 2,  # keep tightly spaced words like "PU HT" apart
                })
                if rows:
                    found.setdefault(number, []).append(pd.DataFrame(rows))
    return found


def extract_text_tables(pdf_path: Path, pages: list[int]) -> dict[int, tuple[str, list[pd.DataFrame]]]:
    results: dict[int, tuple[str, list[pd.DataFrame]]] = {}
    strategies = [
        ("ruled-rows", False, lambda p: _ruled_tables(pdf_path, p)),
        ("camelot-lattice", False, lambda p: _camelot_tables(pdf_path, p, "lattice")),
        ("pdfplumber", False, lambda p: _pdfplumber_tables(pdf_path, p)),
        ("camelot-stream", True, lambda p: _camelot_tables(pdf_path, p, "stream")),
    ]
    remaining = list(pages)
    for method, borderless, extract in strategies:
        if not remaining:
            break
        for number, frames in extract(remaining).items():
            cleaned = [df for df in (clean_table(f, borderless) for f in frames) if is_real_table(df)]
            if cleaned:
                results[number] = (method, cleaned)
        remaining = [p for p in remaining if p not in results]
    return results


# --------------------------------------------------------------------------- #
# 2b. Scanned pages: OCR
# --------------------------------------------------------------------------- #
def extract_ocr_tables(pdf_path: Path, pages: list[int]) -> dict[int, tuple[str, list[pd.DataFrame]]]:
    from img2table.document import PDF
    from img2table.ocr import TesseractOCR

    ocr = TesseractOCR(n_threads=os.cpu_count() or 1, lang=OCR_LANG)
    doc = PDF(str(pdf_path), pages=[p - 1 for p in pages], pdf_text_extraction=False)
    extracted = doc.extract_tables(ocr=ocr, implicit_rows=True, borderless_tables=True, min_confidence=50)

    results: dict[int, tuple[str, list[pd.DataFrame]]] = {}
    for page_index, tables in extracted.items():
        cleaned = [df for df in (clean_table(t.df, borderless=True) for t in tables) if is_real_table(df)]
        if cleaned:
            results[page_index + 1] = ("ocr-table", cleaned)
    return results


# --------------------------------------------------------------------------- #
# 3. Clean the data
# --------------------------------------------------------------------------- #
_WHITESPACE = re.compile(r"\s+")
_NUMBER_TEXT = re.compile(
    r"^(?P<open>\()?(?P<sign>-)?(?P<cur>[$€£¥])?\s?"
    r"(?P<int>\d{1,3}(?:,\d{3})+|\d+)(?:\.(?P<dec>\d+))?(?P<pct>%)?(?P<close>\))?$"
)


def _normalize_cell(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return _WHITESPACE.sub(" ", str(value)).strip()


def excel_number(text: str):
    """Return (value, excel_number_format) so the cell displays exactly like
    `text`, or None when the text must stay text (codes, IDs, words)."""
    match = _NUMBER_TEXT.match(text)
    if not match or bool(match["open"]) != bool(match["close"]) or (match["open"] and match["sign"]):
        return None
    digits, decimals = match["int"].replace(",", ""), match["dec"] or ""
    if len(digits) > 1 and digits.startswith("0"):  # leading zeros: a code, not a number
        return None
    if len(digits) + len(decimals) > 15:  # beyond Excel's precision: long IDs
        return None

    value = float(f"{digits}.{decimals or 0}")
    pattern = ("#,##0" if "," in match["int"] else "0") + (f".{'0' * len(decimals)}" if decimals else "")
    if match["pct"]:
        value /= 100
        pattern += "%"
    if match["cur"]:
        pattern = f'"{match["cur"]}"{pattern}'

    negative = bool(match["open"] or match["sign"])
    if negative and value == 0:
        return None
    if match["open"]:
        pattern = f"{pattern};({pattern})"
    if negative:
        value = -value
    if not decimals and not match["pct"]:
        value = int(value)
    return value, pattern


def _looks_numeric(text) -> bool:
    return isinstance(text, str) and excel_number(text) is not None


def _drop_captions(df: pd.DataFrame) -> pd.DataFrame:
    """Remove title / footnote rows (a single filled cell) above and below the table body."""
    if df.shape[1] < 2:
        return df
    body = [i for i, count in enumerate(df.notna().sum(axis=1)) if count >= 2]
    if body:
        df = df.iloc[body[0] : body[-1] + 1].dropna(axis=1, how="all")
    return df


def _promote_header(df: pd.DataFrame) -> pd.DataFrame:
    first = df.iloc[0]
    if (
        len(df) >= 2
        and first.notna().all()
        and first.nunique() == len(first)
        and not any(_looks_numeric(v) for v in first)
    ):
        df = df.iloc[1:].reset_index(drop=True)
        df.columns = [str(v) for v in first]
    else:
        df.columns = [f"Column {i}" for i in range(1, df.shape[1] + 1)]
    return df


def clean_table(df: pd.DataFrame, borderless: bool = False) -> pd.DataFrame:
    """Keep the table's cells exactly as written; only whitespace inside a cell
    is normalized and fully empty rows/columns are removed. Borderless tables
    can pick up a title or footnote line, which is stripped."""
    if df is None or df.empty:
        return pd.DataFrame()
    df = df.map(_normalize_cell).replace("", pd.NA)
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return pd.DataFrame()
    if borderless:
        df = _drop_captions(df)
    df = df.reset_index(drop=True)
    df.columns = range(df.shape[1])
    return _promote_header(df)


def is_real_table(df: pd.DataFrame) -> bool:
    return df.shape[0] >= 1 and df.shape[1] >= 2


def has_header(df: pd.DataFrame) -> bool:
    """False when the table had no header row (columns are placeholder names)."""
    return list(df.columns) != [f"Column {i}" for i in range(1, df.shape[1] + 1)]


def merge_continued_tables(tables: list[TableResult]) -> list[TableBlock]:
    """Join a table that continues on the next page (same header repeated,
    or same columns without a header) into one block, in PDF order."""
    blocks: list[TableBlock] = []
    for table in tables:
        prev = blocks[-1] if blocks else None
        continues = (
            prev is not None
            and table.page == prev.pages[-1] + 1
            and table.df.shape[1] == prev.df.shape[1]
            and (list(table.df.columns) == list(prev.df.columns) or not has_header(table.df))
        )
        if continues:
            prev.df = pd.concat(
                [prev.df, table.df.set_axis(prev.df.columns, axis=1)], ignore_index=True
            )
            prev.pages.append(table.page)
            prev.methods.append(table.method)
        else:
            blocks.append(TableBlock([table.page], [table.method], table.df.copy()))
    return blocks


# --------------------------------------------------------------------------- #
# 4. Pandas -> Excel (everything on one sheet)
# --------------------------------------------------------------------------- #
def write_excel(blocks: list[TableBlock], out_path: Path) -> None:
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        if not blocks:
            pd.DataFrame({"Result": ["No tables were detected in this PDF."]}).to_excel(
                writer, sheet_name=SHEET_NAME, index=False
            )

        row = 0
        for block in blocks:
            header = has_header(block.df)
            block.df.to_excel(writer, sheet_name=SHEET_NAME, index=False, header=header, startrow=row)
            block.start_row = row + 1
            row += len(block.df) + int(header) + 1  # one blank row between tables

        sheet = writer.sheets[SHEET_NAME]
        if len(blocks) == 1 and has_header(blocks[0].df):
            sheet.freeze_panes = "A2"
        for idx, column in enumerate(sheet.columns, start=1):
            longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)

        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str) and (number := excel_number(cell.value)):
                    cell.value, cell.number_format = number


# --------------------------------------------------------------------------- #
# Full pipeline
# --------------------------------------------------------------------------- #
def convert_pdf(pdf_path: Path, out_path: Path, main_table_only: bool = True) -> ConversionResult:
    """main_table_only keeps the largest table (with its continuation on later
    pages) and drops small side tables such as totals boxes."""
    result = ConversionResult()
    page_has_text = classify_pages(pdf_path)
    text_pages = [p for p, has_text in page_has_text.items() if has_text]
    scanned_pages = [p for p, has_text in page_has_text.items() if not has_text]

    extracted: dict[int, tuple[str, list[pd.DataFrame]]] = {}
    if text_pages:
        extracted.update(extract_text_tables(pdf_path, text_pages))
    if scanned_pages:
        if TESSERACT_AVAILABLE:
            extracted.update(extract_ocr_tables(pdf_path, scanned_pages))
        else:
            result.warnings.append(
                f"{len(scanned_pages)} scanned page(s) skipped: Tesseract OCR is not installed."
            )

    for number, has_text in page_has_text.items():
        method, frames = extracted.get(number, ("none", []))
        if not has_text and not TESSERACT_AVAILABLE:
            method = "skipped (no OCR)"
        result.pages.append(PageReport(number, has_text, method, len(frames)))
        result.tables.extend(TableResult(number, method, df) for df in frames)

    if not result.tables:
        result.warnings.append("No tables were detected in this PDF.")

    result.blocks = merge_continued_tables(result.tables)
    if main_table_only and len(result.blocks) > 1:
        result.blocks = [max(result.blocks, key=lambda b: b.df.size)]
    write_excel(result.blocks, out_path)
    return result
